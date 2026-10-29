#!/usr/bin/env python3
"""Export existing GitHub course HTML to 10 A4 chapter DOCX and one complete DOCX."""
import argparse,json,re,subprocess,shutil
from pathlib import Path
from bs4 import BeautifulSoup,Tag,NavigableString
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def catalog(path):
 s=path.read_text(encoding='utf8')
 d=json.loads(s[s.index('{'):s.rfind('}')+1])
 ch=d['chapters'];lessons=[l for c in ch for l in c['lessons']]
 if len(ch)!=10 or len(lessons)!=120 or [l['number'] for l in lessons]!=list(range(1,121)):
  raise ValueError('Catalog must contain 10 original chapters and 120 ordered lessons')
 return ch,lessons

def clean(x):return re.sub(r'\s+',' ',x.replace('\u00a0',' ')).strip()

def inline(x):
 if isinstance(x,NavigableString):return str(x).replace('\n',' ')
 if not isinstance(x,Tag) or x.name in ('script','style','svg','button','input'):return ''
 t=''.join(inline(c) for c in x.children)
 if x.name in ('strong','b'):return '**'+t.strip()+'**'
 if x.name in ('em','i'):return '*'+t.strip()+'*'
 if x.name=='code' and not x.find_parent('pre'):return chr(96)+t.strip().replace(chr(96),'′')+chr(96)
 if x.name=='a' and x.get('href','').startswith(('http://','https://')):
  return '['+t.strip()+']('+x['href'].replace(' ','%20')+')'
 if x.name=='br':return '  \n'
 if x.name=='sup':return '^'+t.strip()+'^'
 if x.name=='sub':return '~'+t.strip()+'~'
 return t

def blocks(x):
 if not isinstance(x,Tag) or x.name in ('script','style','svg','button','input','nav','header','footer','noscript'):return []
 if x.name in ('h1','h2','h3','h4'):
  return ['#'*{'h1':2,'h2':3,'h3':4,'h4':5}[x.name]+' '+clean(x.get_text(' ',strip=True))]
 if x.name=='pre':
  code=x.find('code')
  src=(code or x).get_text().replace('\r\n','\n').strip('\n')
  lang=[v[9:] for v in (code.get('class',[]) if code else []) if v.startswith('language-')]
  fence=chr(96)*3
  while fence in src:fence+=chr(96)
  return [fence+(lang[0] if lang else '')+'\n'+src+'\n'+fence] if src else []
 if x.name in ('p','div','span') and set(x.get('class',[]))&{'math','math-accent','equation'}:
  s=x.get_text(' ',strip=True)
  s=re.sub(r'^\s*\\\[\s*|\s*\\\]\s*$','',s)
  s=re.sub(r'^\s*\$\$\s*|\s*\$\$\s*$','',s)
  if re.search(r'\\[A-Za-z]+|[_^]',s):return ['$$\n'+s.strip()+'\n$$']
  return ['**数学表达式：** '+s] if s else []
 if x.name=='p':
  s=clean(inline(x))
  return [s] if s else []
 if x.name in ('ul','ol'):
  z=[]
  for i,li in enumerate(x.find_all('li',recursive=False),1):
   z.append((str(i)+'.' if x.name=='ol' else '-')+' '+clean(inline(li)))
  return ['\n'.join(z)] if z else []
 if x.name=='table':
  rows=[]
  for tr in x.find_all('tr'):
   cells=[clean(td.get_text(' ',strip=True)).replace('|','¦') for td in tr.find_all(['th','td'],recursive=False)]
   if cells:rows.append(cells)
  if not rows:return []
  n=max(map(len,rows));rows=[r+['']*(n-len(r)) for r in rows]
  return ['\n'.join(['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*n)+' |']+['| '+' | '.join(r)+' |' for r in rows[1:]])]
 if x.name=='details':
  result=[];s=x.find('summary',recursive=False)
  if s:result.append('**'+clean(s.get_text(' ',strip=True))+'**')
  for y in x.children:
   if isinstance(y,Tag) and y.name!='summary':result+=blocks(y)
  return result
 if x.name=='summary':return ['**'+clean(x.get_text(' ',strip=True))+'**']
 if x.name=='blockquote':return ['> '+clean(x.get_text(' ',strip=True))]
 if x.name=='img':return ['*原网页图示：'+clean(x.get('alt',''))+'*']
 if x.name in ('article','main','section','div','aside','li'):
  result=[]
  for y in x.children:
   if isinstance(y,Tag):result+=blocks(y)
  return result
 return []

def body_from_html(file):
 soup=BeautifulSoup(file.read_text(encoding='utf8'),'html.parser')
 article=soup.find('main') or soup.find('article') or soup.body
 if article is None:raise ValueError('No article/main/body: '+str(file))
 body='\n\n'.join(x for x in blocks(article) if x)
 if len(body)<300:raise ValueError('Incomplete lesson source: '+str(file))
 return body

def reference_doc(file):
 doc=Document();sec=doc.sections[0]
 sec.page_width=Cm(21);sec.page_height=Cm(29.7)
 sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2.3)
 sec.left_margin=Cm(2.55);sec.right_margin=Cm(2.45)
 style=doc.styles['Normal'];style.font.name='Noto Sans CJK SC';style.font.size=Pt(10.5)
 style.font.color.rgb=RGBColor(39,56,69);style.paragraph_format.space_after=Pt(7)
 style.paragraph_format.line_spacing=1.3
 for n,p in [('Title',25),('Heading 1',18),('Heading 2',14),('Heading 3',11)]:
  st=doc.styles[n];st.font.name='Noto Sans CJK SC';st.font.size=Pt(p)
  st.font.bold=True;st.font.color.rgb=RGBColor(25,81,99)
  st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(6)
  st.paragraph_format.keep_with_next=True
 for n in ['Source Code','Verbatim Char']:
  if n in doc.styles:
   doc.styles[n].font.name='Liberation Mono';doc.styles[n].font.size=Pt(8)
 head=sec.header.paragraphs[0];head.alignment=WD_ALIGN_PARAGRAPH.RIGHT
 run=head.add_run('计量经济学120天系统课程 · GitHub正文归档')
 run.font.size=Pt(8);run.font.color.rgb=RGBColor(99,117,126)
 doc.add_paragraph('Reference')
 doc.save(file)

def produce(chapters,lessons,chosen,source,folder,reference,tmp):
 title=('计量经济学120天系统课程｜十章120课全文汇编' if len(chosen)==10 else
        '第'+str(chosen[0])+'章｜'+next(c['title'] for c in chapters if c['number']==chosen[0]))
 filename=('econometrics-all-120-lessons.docx' if len(chosen)==10 else
           'econometrics-chapter-'+f'{chosen[0]:02d}'+'.docx')
 target=folder/filename
 parts=['---','title: '+json.dumps(title,ensure_ascii=False),'lang: zh-CN','---','','# '+title,'',
 '**资料说明**：本Word逐课提取当前GitHub课程HTML现有正文，保留公式、代码和练习答案。所有120课在课程网站上统一显示为“已授课”。这是网站状态设置，不意味着此前100份预编讲义都已经在聊天中逐日正式讲授；未经过真实计算的回归结果不在此虚构。','',
 '**来源**：https://github.com/Ericywh/research-learning-hub/tree/main/docs/econometrics/lessons','',
 '**导出日期**：2026-10-27','']
 headings=[]
 for chapter in chapters:
  if chapter['number'] not in chosen:continue
  parts+=['','\\newpage','','# 第'+str(chapter['number'])+'章｜'+chapter['title'],'']
  for l in chapter['lessons']:
   page=source/'lessons'/(l['id']+'.html')
   if not page.is_file():raise FileNotFoundError(str(page))
   heading=f"第{l['chapter']}章·第{l['section']}节｜{l['title']}"
   label='网站状态：已授课'
   parts+=['','## '+heading,'',f"*全课程第{l['number']}课 / 共120课 · {label}*",'',body_from_html(page),'']
   headings.append(heading)
 md=tmp/(target.stem+'.md');md.write_text('\n'.join(parts),encoding='utf8')
 cmd=['pandoc',str(md),'--from=markdown+tex_math_dollars+tex_math_single_backslash','--to=docx',
      '--standalone','--toc','--toc-depth=2','--reference-doc='+str(reference),'--output='+str(target)]
 p=subprocess.run(cmd,capture_output=True,text=True)
 if p.returncode:raise RuntimeError('Pandoc: '+p.stderr[-4000:])
 if target.stat().st_size<8000:raise RuntimeError('Unexpectedly small DOCX '+str(target))
 doc=Document(target);text='\n'.join(p.text for p in doc.paragraphs)
 if not all(h in text for h in headings):raise RuntimeError('Chapter DOCX missing lesson titles: '+filename)
 return {'filename':filename,'lessons':len(headings),'bytes':target.stat().st_size}

def main():
 p=argparse.ArgumentParser()
 p.add_argument('--source',required=True);p.add_argument('--output',required=True)
 a=p.parse_args();source=Path(a.source);folder=Path(a.output);folder.mkdir(parents=True,exist_ok=True)
 chapters,lessons=catalog(source/'catalog.js')
 tmp=folder/'.build-word';tmp.mkdir(exist_ok=True)
 ref=tmp/'style.docx';reference_doc(ref)
 results=[]
 for chapter in chapters:
  results.append(produce(chapters,lessons,[chapter['number']],source,folder,ref,tmp))
 results.append(produce(chapters,lessons,list(range(1,11)),source,folder,ref,tmp))
 # Actual daily teaching progress comes from the deliberately maintained latest-lesson
 # homepage target, never from the catalog's published display flag.
 landing=(source/'index.html').read_text(encoding='utf8')
 match=re.search(r'<meta[ ]+http-equiv="refresh"[ ]+content="0;[ ]*url=lessons/([0-9]{2}-[0-9]{2})[.]html"',landing,re.I)
 if not match:raise ValueError('Missing authoritative latest-taught lesson redirect in course homepage')
 latest=next((x for x in lessons if x['id']==match.group(1)),None)
 if latest is None:raise ValueError('Homepage points outside the original 120-lesson plan')
 taught=latest['number']
 manifest={'chapters':10,'lessons':120,'taught':taught,'prewritten':120-taught,'documents':results}
 (folder/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf8')
 links=''.join(f'<li><a href="{r["filename"]}">{("完整120课汇编" if r["lessons"]==120 else "分章讲义"+str(i+1))}（{r["lessons"]}课，Word）</a></li>' for i,r in enumerate(results))
 index='<!doctype html><html lang="zh-CN"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>计量经济学120天｜Word下载</title><style>body{font:16px/1.8 system-ui,"Noto Sans SC",sans-serif;max-width:900px;margin:35px auto;padding:20px;color:#264455}a{color:#13777d}li{margin:9px 0}</style><h1>计量经济学120天系统课程｜Word讲义</h1><p>从GitHub课程HTML导出。全部120课在网站上统一显示“已授课”。其中部分正文源自提前编写的教学资料，实际聊天授课历史另行核对。</p><ol>'+links+'</ol><a href="../">返回课程</a></html>'
 (folder/'index.html').write_text(index,encoding='utf8')
 print(json.dumps(manifest,ensure_ascii=False,indent=2))
 shutil.rmtree(tmp)

if __name__=='__main__':main()
